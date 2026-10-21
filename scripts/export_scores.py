#!/usr/bin/env python3
"""Build an auditable score workbook from experiment-report review JSON files.

The only workbook writer. It replaced a Node version that needed ``@oai/artifact-tool``, a package
that ships only inside the Codex runtime -- which made Excel, the one artifact a course actually
collects, impossible on every other machine and agent host. The two were measured equivalent (same
sheets, labels, formulas, number formats, column widths, freeze panes and every cell) before the
Node one was removed, so nothing about the deliverable changed.

What this writer deliberately does not do: render sheet PNG previews, or dump the artifact tool's
formula scan. Neither reached a deliverable even while the Node writer existed (the server deleted
its preview directory immediately, and the scan output was never parsed), and the same arithmetic
is validated here before a file is written -- which fails earlier than a post-export scan can.

Usage: export_scores.py <review-dir> <output.xlsx>
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

YAHEI = "Microsoft YaHei"
THIN = Border(
    left=Side(style="thin", color="FF808080"),
    right=Side(style="thin", color="FF808080"),
    top=Side(style="thin", color="FF808080"),
    bottom=Side(style="thin", color="FF808080"),
)
WHITE = PatternFill(fill_type="solid", fgColor="FFFFFFFF")

SHORT_CRITERION = {
    "实验报告完整性": "完整性",
    "数据记录与计算完整性": "数据与计算",
    "结论合理性": "结论",
    "小结质量": "小结",
    "格式规范性与语言表达": "格式与表达",
}
CRITERION_START_COL = 3


def column_letter(index: int) -> str:
    result = ""
    value = index
    while value > 0:
        value -= 1
        result = chr(65 + (value % 26)) + result
        value //= 26
    return result


def load_reviews(review_dir: Path) -> list[dict]:
    names = sorted(
        name for name in (entry.name for entry in review_dir.iterdir())
        if name.lower().endswith(".review.json")
    )
    if not names:
        raise ValueError("No review JSON files found")
    reviews = []
    for name in names:
        data = json.loads((review_dir / name).read_text(encoding="utf-8"))
        if not data.get("student_name") or not data.get("student_id"):
            raise ValueError(f"{name}: student_name and student_id are required")
        items = data.get("items")
        if not items:
            raise ValueError(f"{name}: items are required")
        total = sum(float(item["score"]) for item in items)
        maximum = sum(float(item["max_score"]) for item in items)
        if abs(total - float(data["total_score"])) > 0.01:
            raise ValueError(f"{name}: item scores do not sum to total_score")
        if abs(maximum - float(data["max_score"])) > 0.01:
            raise ValueError(f"{name}: item maxima do not sum to max_score")
        reviews.append(data)

    rubric = [(item["criterion"], float(item["max_score"])) for item in reviews[0]["items"]]
    for review in reviews[1:]:
        current = [(item["criterion"], float(item["max_score"])) for item in review["items"]]
        if current != rubric:
            raise ValueError("All reviews must use the same rubric before one workbook can be exported")
    return reviews


def style_range(sheet, ref: str, **kwargs) -> None:
    for row in sheet[ref]:
        for cell in row:
            if "font" in kwargs:
                cell.font = kwargs["font"]
            if kwargs.get("fill"):
                cell.fill = WHITE
            if kwargs.get("border"):
                cell.border = THIN
            if kwargs.get("align") is not None:
                cell.alignment = kwargs["align"]
            if "number_format" in kwargs:
                cell.number_format = kwargs["number_format"]
            if kwargs.get("wrap"):
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)


def build(review_dir: Path, output_path: Path) -> dict:
    reviews = load_reviews(review_dir)
    rubric = [(item["criterion"], float(item["max_score"])) for item in reviews[0]["items"]]

    uses_offset = any(float(review.get("score_offset") or 0) > 0 for review in reviews)
    total_col = CRITERION_START_COL + len(rubric)
    adjusted_col = total_col + 1 if uses_offset else None
    max_col = (adjusted_col or total_col) + 1
    status_col = max_col + 1
    source_col = status_col + 1
    last_col = column_letter(source_col)
    first_criterion_letter = column_letter(CRITERION_START_COL)
    last_criterion_letter = column_letter(total_col - 1)
    total_letter = column_letter(total_col)
    adjusted_letter = column_letter(adjusted_col) if adjusted_col else None
    max_letter = column_letter(max_col)
    status_letter = column_letter(status_col)
    source_letter = column_letter(source_col)

    workbook = Workbook()
    summary = workbook.active
    summary.title = "成绩汇总"
    rubric_sheet = workbook.create_sheet("评分标准")
    summary.sheet_view.showGridLines = False
    rubric_sheet.sheet_view.showGridLines = False

    summary.merge_cells(f"A1:{last_col}1")
    summary["A1"] = "实验报告成绩汇总"
    title_font = Font(bold=True, color="FF000000", size=16, name=YAHEI)
    summary["A1"].font = title_font
    summary["A1"].alignment = Alignment(horizontal="left", vertical="center")
    summary.row_dimensions[1].height = 30
    style_range(summary, f"A1:{last_col}1", font=title_font, fill=True, border=True)

    start_row = 7
    end_row = start_row + len(reviews) - 1
    summary["A3"] = "报告数"
    summary["C3"] = "已审核"
    summary["E3"] = "原始平均分（批改得分）" if uses_offset else "平均分"
    summary_block_end = 4 if uses_offset else 3
    block_font = Font(bold=True, color="FF000000", name=YAHEI, size=11)
    for row in summary[f"A3:F{summary_block_end}"]:
        for cell in row:
            cell.font = block_font
            cell.fill = WHITE
            cell.border = THIN
            cell.alignment = Alignment(horizontal="left")

    summary["B3"] = f"=COUNTA(A{start_row}:A{end_row})"
    # Formula text is copied verbatim from the .mjs, including its lack of $ anchors on the
    # summary row, so the equivalence diff compares formulas rather than normalising them.
    summary["D3"] = f'=COUNTIF({status_letter}{start_row}:{status_letter}{end_row},"已审核")'
    summary["F3"] = f"=AVERAGE({total_letter}{start_row}:{total_letter}{end_row})"
    for coordinate in ("B3", "C3", "D3"):
        summary[coordinate].number_format = "0"
    summary["F3"].number_format = "0.0"
    if uses_offset:
        summary["E4"] = "调整后平均分（上报成绩）"
        summary["F4"] = f"=AVERAGE({adjusted_letter}{start_row}:{adjusted_letter}{end_row})"
        summary["F4"].number_format = "0.0"
        for row in summary["A4:F4"]:
            for cell in row:
                cell.font = block_font
                cell.fill = WHITE
                cell.border = THIN
                cell.alignment = Alignment(horizontal="left")

    headers = ["姓名", "学号"]
    headers += [f"{SHORT_CRITERION.get(name, name)} ({int(maximum)})" for name, maximum in rubric]
    headers += ["总分"]
    if uses_offset:
        headers += ["调整后成绩"]
    headers += ["满分", "审核状态", "源文件"]
    for index, header in enumerate(headers, start=1):
        summary.cell(row=6, column=index, value=header)

    for offset, review in enumerate(reviews):
        row = start_row + offset
        summary.cell(row=row, column=1, value=str(review["student_name"]))
        summary.cell(row=row, column=2, value=f"\u2060{review['student_id']}")
        for index, item in enumerate(review["items"]):
            summary.cell(row=row, column=CRITERION_START_COL + index, value=float(item["score"]))
        summary.cell(row=row, column=max_col, value=float(review["max_score"]))
        status = (
            "已审核" if review.get("review_status") == "approved"
            else "待裁决" if review.get("adjudication_status") == "required"
            else "待审核"
        )
        summary.cell(row=row, column=status_col, value=status)
        summary.cell(row=row, column=source_col, value=str(review.get("source_file") or ""))
        summary.cell(
            row=row,
            column=total_col,
            value=f"=SUM({first_criterion_letter}{row}:{last_criterion_letter}{row})",
        )
        if uses_offset:
            review_offset = float(review.get("score_offset") or 0)
            summary.cell(
                row=row,
                column=adjusted_col,
                value=f"=MIN({max_letter}{row},{total_letter}{row}+{review_offset:g})",
            )

    data_ref = f"A6:{last_col}{end_row}"
    style_range(
        summary,
        data_ref,
        font=Font(name=YAHEI, size=10),
        fill=True,
        border=True,
        align=Alignment(horizontal="left", vertical="center"),
    )
    style_range(summary, f"A6:{last_col}6", font=Font(bold=True, color="FF000000", name=YAHEI, size=10))
    style_range(summary, f"A{start_row}:{last_col}{end_row}", font=Font(name=YAHEI, size=10), wrap=True)
    style_range(summary, f"B{start_row}:B{end_row}", font=Font(name=YAHEI, size=10), number_format="@")
    style_range(
        summary,
        f"{first_criterion_letter}{start_row}:{max_letter}{end_row}",
        font=Font(name=YAHEI, size=10),
        number_format="0.0",
    )
    summary.freeze_panes = f"A{start_row}"
    # Letters are derived, never written out: without a score offset the adjusted column does
    # not exist, so a hardcoded status/source letter would land on the wrong columns. (It did:
    # the panel and the workbook disagreed for every batch graded without an offset.)
    widths = {"A": 14, "B": 18, status_letter: 12, source_letter: 46}
    for column in range(CRITERION_START_COL, total_col):
        widths[column_letter(column)] = 17
    for column in range(total_col, max_col + 1):
        widths[column_letter(column)] = 10
    for letter, width in widths.items():
        summary.column_dimensions[letter].width = width
    summary.add_table(
        Table(
            displayName="ScoreSummaryTable",
            ref=data_ref,
            tableStyleInfo=TableStyleInfo(name="TableStyleLight1", showRowStripes=False),
        )
    )

    rubric_end = 3 + len(rubric)
    rubric_sheet.merge_cells("A1:C1")
    rubric_sheet["A1"] = "本批次确认的评分标准"
    style_range(
        rubric_sheet,
        "A1:C1",
        font=Font(bold=True, color="FF000000", size=15, name=YAHEI),
        fill=True,
        border=True,
    )
    rubric_sheet["A1"].alignment = Alignment(horizontal="left")
    for column, header in enumerate(["评分维度", "分值", "权重"], start=1):
        rubric_sheet.cell(row=3, column=column, value=header)
    for offset, (criterion, maximum) in enumerate(rubric):
        row = 4 + offset
        rubric_sheet.cell(row=row, column=1, value=criterion)
        rubric_sheet.cell(row=row, column=2, value=maximum)
        rubric_sheet.cell(row=row, column=3, value=f"=B{row}/SUM($B$4:$B${rubric_end})")
    rubric_ref = f"A3:C{rubric_end}"
    style_range(
        rubric_sheet,
        rubric_ref,
        font=Font(name=YAHEI, size=10),
        fill=True,
        border=True,
        align=Alignment(horizontal="left"),
    )
    style_range(rubric_sheet, "A3:C3", font=Font(bold=True, color="FF000000", name=YAHEI, size=10))
    style_range(rubric_sheet, f"A4:A{rubric_end}", font=Font(name=YAHEI, size=10), wrap=True)
    for row in rubric_sheet[f"A4:A{rubric_end}"]:
        for cell in row:
            cell.alignment = Alignment(horizontal="left", wrap_text=True)
    style_range(rubric_sheet, f"B4:B{rubric_end}", font=Font(name=YAHEI, size=10), number_format="0.0")
    style_range(rubric_sheet, f"C4:C{rubric_end}", font=Font(name=YAHEI, size=10), number_format="0.0%")
    rubric_sheet.freeze_panes = "A4"
    rubric_sheet.column_dimensions["A"].width = 34
    rubric_sheet.column_dimensions["B"].width = 14
    rubric_sheet.column_dimensions["C"].width = 14
    rubric_sheet.add_table(
        Table(
            displayName="RubricTable",
            ref=rubric_ref,
            tableStyleInfo=TableStyleInfo(name="TableStyleLight1", showRowStripes=False),
        )
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return {"output": str(output_path), "rows": len(reviews), "sheets": workbook.sheetnames}


def snapshot_workbook(path: Path) -> dict:
    """Everything a teacher could see, reduced to comparable form.

    ``openpyxl`` reads both the artifact-tool workbook and this writer's workbook, which is what
    makes an equivalence check possible without opening Excel.
    """
    workbook = openpyxl.load_workbook(path, data_only=False)
    sheets: dict[str, dict] = {}
    for name in workbook.sheetnames:
        sheet = workbook[name]
        sheets[name] = {
            "freeze_panes": sheet.freeze_panes,
            "show_grid_lines": sheet.sheet_view.showGridLines,
            "merged": sorted(str(item) for item in sheet.merged_cells.ranges),
            "column_widths": {
                key: dimension.width
                for key, dimension in sheet.column_dimensions.items()
                if dimension.width is not None
            },
            "row_heights": {
                key: dimension.height
                for key, dimension in sheet.row_dimensions.items()
                if dimension.height is not None
            },
            "tables": {key: str(getattr(value, "ref", value)) for key, value in sheet.tables.items()},
            "cells": {
                cell.coordinate: {
                    "value": cell.value,
                    "number_format": cell.number_format,
                    "font": (cell.font.name, cell.font.size, bool(cell.font.bold)),
                    "align": (
                        None if cell.alignment is None else cell.alignment.horizontal,
                        None if cell.alignment is None else cell.alignment.vertical,
                        None if cell.alignment is None else bool(cell.alignment.wrap_text),
                    ),
                }
                for row in sheet.iter_rows()
                for cell in row
                if cell.value is not None
            },
        }
    return {"sheets": sheets}


def compare_workbooks(left_path: Path, right_path: Path) -> list[str]:
    """Differences between two workbooks, empty when they are equivalent."""
    left, right = snapshot_workbook(left_path), snapshot_workbook(right_path)
    diffs: list[str] = []
    if list(left["sheets"]) != list(right["sheets"]):
        diffs.append(f"sheet names differ: {list(left['sheets'])} vs {list(right['sheets'])}")
    for name in left["sheets"]:
        if name not in right["sheets"]:
            continue
        a, b = left["sheets"][name], right["sheets"][name]
        for field in ("freeze_panes", "show_grid_lines", "merged", "column_widths", "row_heights", "tables"):
            if a[field] != b[field]:
                diffs.append(f"[{name}] {field}: {a[field]} != {b[field]}")
        for coordinate in sorted(set(a["cells"]) | set(b["cells"])):
            if a["cells"].get(coordinate) != b["cells"].get(coordinate):
                diffs.append(
                    f"[{name}] {coordinate}: {a['cells'].get(coordinate)} != {b['cells'].get(coordinate)}"
                )
    return diffs


def main() -> int:
    if len(sys.argv) == 4 and sys.argv[1] == "--compare":
        diffs = compare_workbooks(Path(sys.argv[2]).resolve(), Path(sys.argv[3]).resolve())
        print(json.dumps({"identical": not diffs, "differences": len(diffs)}, ensure_ascii=False))
        for item in diffs[:20]:
            print(item)
        return 0 if not diffs else 1
    if len(sys.argv) != 3:
        print(
            "Usage: export_scores.py <review-dir> <output.xlsx>\n"
            "       export_scores.py --compare <a.xlsx> <b.xlsx>",
            file=sys.stderr,
        )
        return 2
    try:
        result = build(Path(sys.argv[1]).resolve(), Path(sys.argv[2]).resolve())
    except ValueError as error:
        print(json.dumps({"ok": False, "error": str(error)}, ensure_ascii=False), file=sys.stderr)
        return 1
    print(json.dumps({**result, "ok": True}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
